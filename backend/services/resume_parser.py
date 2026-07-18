import os
import re
from typing import Dict, List, Any
import fitz  # PyMuPDF
import docx
from backend.schemas.resume import ResumeParsedData

def parse_pdf(file_path: str) -> str:
    text = ""
    try:
        doc = fitz.open(file_path)
        for page in doc:
            text += page.get_text()
        doc.close()
    except Exception as e:
        print(f"Error parsing PDF: {e}")
    return text

def parse_docx(file_path: str) -> str:
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        print(f"Error parsing DOCX: {e}")
    return text

def extract_skills(text: str) -> List[str]:
    # Basic keyword extraction - in a real app, this would use NLP or an AI API
    common_skills = [
        "python", "java", "javascript", "c++", "c#", "ruby", "go", "rust",
        "react", "angular", "vue", "node.js", "express", "django", "fastapi", "flask",
        "sql", "mysql", "postgresql", "mongodb", "redis", "elasticsearch",
        "aws", "azure", "gcp", "docker", "kubernetes", "ci/cd", "git",
        "machine learning", "data science", "ai", "nlp", "computer vision"
    ]
    
    text_lower = text.lower()
    found_skills = set()
    
    for skill in common_skills:
        # Simple word boundary regex
        pattern = r'\b' + re.escape(skill) + r'\b'
        if re.search(pattern, text_lower):
            found_skills.add(skill)
            
    return list(found_skills)

def extract_sections(text: str) -> Dict[str, Any]:
    # Very basic section extraction - in a real app, use AI or more complex parsing
    # Here we just return placeholders since full extraction requires complex logic
    return {
        "experience": [{"company": "Extracted Company", "title": "Extracted Title", "duration": "", "bullets": []}],
        "education": [{"school": "Extracted School", "degree": "Extracted Degree", "field": "", "year": ""}],
        "skills": [],
        "certifications": []
    }

def parse_resume(file_path: str, file_type: str) -> ResumeParsedData:
    """Dispatcher for parsing different resume file types."""
    if file_type == "pdf":
        raw_text = parse_pdf(file_path)
    elif file_type == "docx":
        raw_text = parse_docx(file_path)
    else:
        raw_text = ""
        
    skills = extract_skills(raw_text)
    sections = extract_sections(raw_text)
    
    # Merge skills
    all_skills = list(set(skills + sections.get("skills", [])))
    
    return ResumeParsedData(
        skills=all_skills,
        experience=sections.get("experience", []),
        education=sections.get("education", []),
        certifications=sections.get("certifications", []),
        raw_text=raw_text
    )
